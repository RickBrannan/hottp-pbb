"""
    convert hottp.xml into two different PBB files,
    only English for now
    RWB 2023-12-08
"""
import os
import regex as re
from hebrew import Hebrew
from docx import Document
from unicodedata import normalize
from biblelib.word import BCVWPID, BCVID, fromusfm
from biblelib.book import book
from lxml import etree
from src import DATAPATH, ROOT


hottp_dir = ROOT.parent.parent / "ubsicap/ubs-open-license/HOTTP/"
data_dir = DATAPATH / "pbb/"
book_names = book.Books()
current_book = "Genesis"

lang = "en"
# lang = "fr"


def get_current_book(entry):
    for reference in entry.xpath(".//Reference"):
        marble_id = reference.text
        bcv = BCVID(re.sub(r"^0(\d{8})\d{5}$",r"\1", marble_id))
        current_book = book_names.fromusfmnumber(bcv.book_ID).name
        return current_book
    return ""


def get_current_reference(entry):
    refs = []
    for reference in entry.xpath(".//Reference"):
        marble_id = reference.text
        bcv = BCVID(re.sub(r"^0(\d{8})\d{5}$",r"\1", marble_id))
        if not bcv.ID in refs:
            refs.append(bcv.ID)

    if len(refs) == 1:
        return BCVID(refs[0]).to_usfm()
    elif len(refs) == 0:
        # we have problems. re-get the first reference and roll with it.
        marble_id = entry.xpath(".//Reference")[0].text
        bcv = BCVID(re.sub(r"^0(\d{8})\d{5}$", r"\1", marble_id))
        return bcv.to_usfm()
    else:
        return f"{BCVID(refs[0]).to_usfm()}–{int(BCVID(refs[-1]).verse_ID)}"


# get a word doc started since that's what PBB uses
pbb_doc = Document()

# title page? maybe later

# open HOTTP.XML and start parsin'.
hottp_xml = etree.parse(hottp_dir / "hottp.xml")
hottp_root = hottp_xml.getroot()
previous_book = "Genesis"
previous_reference = ""
pbb_doc.add_heading("Genesis", level=1)
for entry in hottp_root.xpath(".//HOTTP_Entry"):
    id = entry.attrib["Id"]
    current_reference = get_current_reference(entry)
    current_book = get_current_book(entry)
    if current_book != previous_book:
        previous_book = current_book
        pbb_doc.add_heading(current_book, level=1)

    print(f"{id}: {current_reference}")
    if current_reference != previous_reference:
        pbb_doc.add_heading(f"[[@BibleBHS:{current_reference}]][[BibleBHS:{current_reference}]]", level=2)
        previous_reference = current_reference
    else:
        pbb_doc.add_heading(f"[[BibleBHS:{current_reference}]]", level=2)

    remark = ""
    if lang == "en":
        remark = entry.xpath(".//Remark")[0].text
    elif lang == "fr":
        remark = entry.xpath(".//Remark_FR")[0].text
    suggestion = ""
    if lang == "en":
        suggestion = entry.xpath(".//Suggestion")[0].text
    elif lang == "fr":
        suggestion = entry.xpath(".//SuggestionFR")[0].text

    if remark != "":
        remark_paragraph = pbb_doc.add_paragraph()
        # need to localize for FR
        remark_paragraph.add_run("Remark:").bold = True
        remark_paragraph.add_run(" ")
        # because it looks like there are at least &lt;span...&gt; in some elements
        # this is usually Hebrew; need to actually treat it properly at some point.
        remark = re.sub(r"<[^>]+>","", str(remark))
        remark_paragraph.add_run(remark)
    if suggestion != "":
        suggestion_paragraph = pbb_doc.add_paragraph()
        # need to localize for FR
        suggestion_paragraph.add_run("Suggestion:").bold = True
        suggestion_paragraph.add_run(" ")
        suggestion_paragraph.add_run(suggestion)

    # cycle the alternatives
    alt_count = 0
    for alternative in entry.xpath(".//Alternatives/Alternative"):
        # Source, Rating, Versions/Version/Text, Versions/Version/Content, Factors, Literal, LiteralFR
        alt_count += 1
        pbb_doc.add_heading(f"Alternative {alt_count}", level=3)
        # need to get a hebrew text style around "Source"
        source = alternative.xpath(".//Source")[0].text
        pbb_doc.add_paragraph(f"{source}")
        rating = alternative.xpath(".//Rating")[0].text
        # need to figure out how to reference Rating and Factors in HTML
        pbb_doc.add_paragraph(f"Rating: {rating}")

        # need to map versions to languages and get proper langauges in here
        for version in alternative.xpath(".//Versions/Version"):
            version_text = version.xpath(".//Text")[0].text
            version_content = version.xpath(".//Content")[0].text
            version_paragraph = pbb_doc.add_paragraph(f"{version_text}:", style="List Bullet")
            version_paragraph.add_run(f" {version_content}").italic = True

        factors = []
        for factor in alternative.xpath(".//Factors/Factor"):
            factors.append(factor.text)
        if len(factors) > 0:
            pbb_doc.add_paragraph(f"Factors: {', '.join(factors)}")

        if lang == "en":
            literal = alternative.xpath(".//Literal")[0].text
            literal_paragraph = pbb_doc.add_paragraph()
            literal_paragraph.add_run(f"Literal: ").bold = True
            literal_paragraph.add_run(f"{literal}")
        elif lang == "fr":
            literal = alternative.xpath(".//LiteralFR")[0].text
            literal_paragraph = pbb_doc.add_paragraph()
            literal_paragraph.add_run(f"Literal: ").bold = True
            literal_paragraph.add_run(f"{literal}")

pbb_doc.save(data_dir / f"hottp_{lang}.docx")


